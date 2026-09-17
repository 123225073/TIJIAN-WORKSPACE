import io
import pytest
from docx import Document
from pypdf import PdfWriter
from backend.documents import extract

def test_docx_extracts_paragraphs_and_tables():
    doc=Document();doc.add_paragraph('电梯更新调研');table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='区域';table.cell(0,1).text='上海'
    data=io.BytesIO();doc.save(data)
    text=extract('调研.docx',data.getvalue())
    assert '电梯更新调研' in text and '区域 | 上海' in text

def test_empty_pdf_and_unsupported_document_rejected():
    pdf=PdfWriter();pdf.add_blank_page(width=200,height=200);data=io.BytesIO();pdf.write(data)
    with pytest.raises(ValueError,match='OCR'):extract('扫描.pdf',data.getvalue())
    with pytest.raises(ValueError):extract('执行.exe',b'binary')
