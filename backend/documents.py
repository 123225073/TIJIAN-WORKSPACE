"""Bounded, local document text extraction. Never execute embedded content."""
import io,zipfile
from pathlib import Path

def extract(filename,raw):
    suffix=Path(filename).suffix.lower()
    if len(raw)>10_000_000:raise ValueError('资料导入上限10MB')
    if suffix=='.pdf':
        from pypdf import PdfReader
        try:
            pdf=PdfReader(io.BytesIO(raw))
            if pdf.is_encrypted:raise ValueError('请先在本机解密PDF，再导入')
            if len(pdf.pages)>200:raise ValueError('PDF最多200页，请拆分后导入')
            pages=[p.extract_text() or '' for p in pdf.pages]
            text='\n\n'.join(f'【第{i+1}页】\n'+p for i,p in enumerate(pages))
            if len(''.join(pages).strip())<30:raise ValueError('未读取到有效文字，扫描PDF请先OCR或导入文字稿')
        except ValueError:raise
        except Exception:raise ValueError('PDF无法读取，请检查文件是否完整')
    elif suffix=='.docx':
        from docx import Document
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(x.file_size for x in z.infolist())>40_000_000:raise ValueError('文档展开后超过40MB，请拆分后导入')
            doc=Document(io.BytesIO(raw))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        except ValueError:raise
        except Exception:raise ValueError('Word文档无法读取，请另存为DOCX后导入')
    elif suffix in ['.md','.txt','.srt','.vtt','.csv']:
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:raise ValueError('请将文稿保存为UTF-8编码')
    else:raise ValueError('支持PDF、DOCX、Markdown、文本、字幕或CSV')
    if len(text)>500_000:raise ValueError('提取文字超过50万字符，请拆分资料')
    if not text.strip():raise ValueError('文件未包含可读取正文')
    return text
